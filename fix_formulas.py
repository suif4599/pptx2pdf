"""Replace Equation.3 / Equation.DSMT4 OLE objects in a pptx with images
rendered from their WMF/EMF previews, and retarget standalone p:pic shapes
backed by WMF/EMF media the same way, so viewers without OLE/metafile support
(genoffice, LibreOffice PDF export) still show them.

By default EMF-derived images embed the converter's SVG through the
asvg:svgBlip extension with a blank PNG as the raster placeholder (the PDF
export honors the SVG, so rasterizing the fallback would be wasted work).
With --nosvg everything is rasterized by inkscape instead (formula WMF via
metafile-rs at 600 dpi, page EMF via libemf2svg at 300 dpi).
Run inside the flake font environment."""
import hashlib
import os
import subprocess
import sys
import time

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu
from tqdm import tqdm

Image.MAX_IMAGE_PIXELS = None

FORMULA_DPI = 600
PAGE_DPI = 300
MAX_PX = 8000
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
A14_M = "{http://schemas.microsoft.com/office/drawing/2010/main}m"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


PNG_DIR = "."
EMBED_SVG = True


def blank_png():
    placeholder = os.path.join(PNG_DIR, "_blank.png")
    if not os.path.exists(placeholder):
        Image.new("RGBA", (2, 2), (0, 0, 0, 0)).save(placeholder)
    return placeholder


def convert_media(key, blob, ext, dpi):
    svg = os.path.join(PNG_DIR, key + ".svg")
    if EMBED_SVG and ext == ".emf":
        if not os.path.exists(svg):
            src = os.path.join(PNG_DIR, key + ext)
            raw = os.path.join(PNG_DIR, key + "_raw.svg")
            with open(src, "wb") as f:
                f.write(blob)
            try:
                subprocess.run(["emf2svg-conv", "-p", "-i", src, "-o", raw],
                               check=True, capture_output=True, timeout=300)
                # LO draws svgBlip <text> as live text, overlapping the
                # formulas rendered on top — flatten glyphs to paths
                subprocess.run(["inkscape", raw, "--export-text-to-path",
                                "--export-plain-svg", "-o", svg],
                               check=True, capture_output=True, timeout=300)
            finally:
                os.remove(src)
                if os.path.exists(raw):
                    os.remove(raw)
        return blank_png(), svg
    png = os.path.join(PNG_DIR, key + ".png")
    if os.path.exists(png):
        return png, None
    src = os.path.join(PNG_DIR, key + ext)
    with open(src, "wb") as f:
        f.write(blob)
    to_svg = (
        ["emf2svg-conv", "-p", "-i", src, "-o", svg]
        if ext == ".emf"
        else ["mf-cli", src, svg]
    )

    def rasterize(d):
        subprocess.run(
            ["inkscape", svg, "--export-type=png", "--export-filename", png, "-d", str(d)],
            check=True, capture_output=True, timeout=300,
        )

    try:
        subprocess.run(to_svg, check=True, capture_output=True, timeout=300)
        rasterize(dpi)
        with Image.open(png) as im:
            longest = max(im.size)
        if longest > MAX_PX:
            rasterize(max(72, dpi * MAX_PX // longest))
    finally:
        os.remove(src)
        if os.path.exists(svg):
            os.remove(svg)
    return png, None


svg_parts = {}

def attach_svg(slide, blip, svg_path):
    # media shared across slides must map to one part or the zip gets duplicate entries
    part = svg_parts.get(svg_path)
    if part is None:
        with open(svg_path, "rb") as f:
            part = Part(PackURI(f"/ppt/media/{os.path.basename(svg_path)}"), "image/svg+xml", slide.part.package, f.read())
        svg_parts[svg_path] = part
    rid = slide.part.relate_to(part, RT.IMAGE)
    ext_lst = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", SVG_EXT_URI)
    svg_blip = etree.SubElement(ext, "{%s}svgBlip" % ASVG_NS)
    svg_blip.set(qn("r:embed"), rid)


def media_key(slide, rid):
    part = slide.part.related_part(rid)
    name = os.path.basename(str(part.partname))
    key, extn = name.rsplit(".", 1)
    return key, extn.lower(), part.blob


def unwrap_math_fallback(slide):
    # LO renders both branches of a14 math AlternateContent, ghosting the
    # formula over its own image; keep only the image-bearing Fallback
    n = 0
    for ac in slide.shapes._spTree.findall(f".//{MC}AlternateContent"):
        choice = ac.find(MC + "Choice")
        fallback = ac.find(MC + "Fallback")
        if choice is None or fallback is None or choice.find(f".//{A14_M}") is None:
            continue
        if not fallback.findall(f".//{{{NS['a']}}}blip"):
            continue
        parent = ac.getparent()
        idx = parent.index(ac)
        for child in list(fallback):
            idx += 1
            parent.insert(idx, child)
        parent.remove(ac)
        n += 1
    return n


def main():
    global PNG_DIR, EMBED_SVG
    flags = [a for a in sys.argv[1:] if a.startswith("--")]
    positional = [a for a in sys.argv[1:] if not a.startswith("--")]
    EMBED_SVG = "--nosvg" not in flags
    if len(positional) != 2:
        print("usage: fix_formulas.py [--nosvg] <input.pptx> <output.pptx>", file=sys.stderr)
        return 1
    src, out = positional
    digest = hashlib.sha1(open(src, "rb").read()).hexdigest()[:16]
    cache_root = os.environ.get("XDG_CACHE_HOME") or os.path.expanduser("~/.cache")
    PNG_DIR = os.path.join(cache_root, "pptx2pdf", digest)

    os.makedirs(PNG_DIR, exist_ok=True)
    prs = Presentation(src)

    unwrapped = sum(unwrap_math_fallback(s) for s in prs.slides)
    print(f"math fallbacks unwrapped: {unwrapped}")

    jobs = []
    pic_jobs = []
    media = {}
    for s_idx, slide in enumerate(prs.slides, 1):
        for frame in slide.shapes._spTree.xpath(".//p:graphicFrame"):
            prog_ids = frame.xpath(".//p:oleObj/@progId")
            if not any(pid.startswith("Equation") for pid in prog_ids):
                continue
            blips = frame.xpath(".//a:blip/@r:embed")
            xfrm = frame.find("p:xfrm", NS)
            off = xfrm.find("a:off", NS)
            ext = xfrm.find("a:ext", NS)
            if not blips:
                print(f"slide {s_idx}: equation frame without preview blip, skipped")
                continue
            key, extn, blob = media_key(slide, blips[0])
            media.setdefault(key, (blob, "." + extn, FORMULA_DPI))
            jobs.append((s_idx, slide, frame, key,
                         int(off.get("x")), int(off.get("y")),
                         int(ext.get("cx")), int(ext.get("cy"))))
        for el in slide.shapes._spTree.findall(f".//{{{NS['p']}}}pic"):
            if el.xpath("ancestor::p:graphicFrame"):
                continue
            blips = el.xpath(".//a:blip/@r:embed")
            if not blips:
                continue
            key, extn, blob = media_key(slide, blips[0])
            if extn not in ("wmf", "emf"):
                continue
            media.setdefault(key, (blob, "." + extn, PAGE_DPI))
            pic_jobs.append((slide, el, blips[0], key))

    used = {j[3] for j in jobs} | {j[3] for j in pic_jobs}
    media = {k: v for k, v in media.items() if k in used}

    print(f"equation frames: {len(jobs)}, standalone metafile pics: {len(pic_jobs)}, "
          f"unique media: {len(media)}")

    pngs = {}
    svgs = {}
    with tqdm(media.items(), total=len(media), desc="rendering", unit="file") as pbar:
        for key, (blob, extn, dpi) in pbar:
            t0 = time.monotonic()
            pngs[key], svgs[key] = convert_media(key, blob, extn, dpi)
            pbar.set_postfix_str(f"{key}{extn} {time.monotonic() - t0:.1f}s")

    # a full-page EMF snapshot already carries the slide's formulas; a
    # replacement stacked on them would double-render, so skip those frames
    snaps = {}
    for s_idx, slide in enumerate(prs.slides, 1):
        for el in slide.shapes._spTree.findall(f".//{{{NS['p']}}}pic"):
            blips = el.xpath(".//a:blip/@r:embed")
            if not blips:
                continue
            _, extn, _ = media_key(slide, blips[0])
            xf = el.find("p:spPr/a:xfrm", NS)
            if extn != "emf" or xf is None:
                continue
            off, ext = xf.find("a:off", NS), xf.find("a:ext", NS)
            if (int(ext.get("cx")) >= 0.9 * prs.slide_width
                    and int(ext.get("cy")) >= 0.9 * prs.slide_height):
                snaps.setdefault(s_idx, []).append(
                    (int(off.get("x")), int(off.get("y")),
                     int(off.get("x")) + int(ext.get("cx")),
                     int(off.get("y")) + int(ext.get("cy"))))

    replaced = skipped = 0
    for s_idx, slide, frame, key, x, y, cx, cy in jobs:
        parent = frame.getparent()
        idx = parent.index(frame)
        parent.remove(frame)
        if any(x >= sx0 and y >= sy0 and x + cx <= sx1 and y + cy <= sy1
               for sx0, sy0, sx1, sy1 in snaps.get(s_idx, ())):
            skipped += 1
            continue
        pic = slide.shapes.add_picture(pngs[key], Emu(x), Emu(y), Emu(cx), Emu(cy))
        parent.insert(idx, pic._element)
        if svgs[key]:
            attach_svg(slide, pic._element.xpath(".//a:blip")[0], svgs[key])
        replaced += 1
    print(f"replaced: {replaced} frames, skipped on snapshots: {skipped}")

    for slide, el, rid, key in pic_jobs:
        blip = el.xpath(".//a:blip")[0]
        _, new_rid = slide.part.get_or_add_image_part(pngs[key])
        blip.set(qn("r:embed"), new_rid)
        if svgs[key]:
            attach_svg(slide, blip, svgs[key])
    print(f"retargeted: {len(pic_jobs)} pics")

    # the ppt-era converter leaves each OLE visual duplicated in an orphan
    # vmlDrawing part; LO renders those WMF shapes even with the frame gone,
    # ghosting the formula over its replacement — drop legacy rels once no
    # oleObj references them anymore
    dropped = 0
    for slide in prs.slides:
        if slide.shapes._spTree.xpath(".//p:oleObj"):
            continue
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype.endswith("/oleObject") or rel.reltype.endswith("/vmlDrawing"):
                slide.part.drop_rel(rid)
                dropped += 1
    print(f"legacy ole/vml rels dropped: {dropped}")

    prs.save(out)
    print(f"saved: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
